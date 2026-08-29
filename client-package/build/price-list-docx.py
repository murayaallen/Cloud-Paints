# ============================================================
# CLOUD PAINTS — the price list as an editable Word document
# ============================================================
#     python build/price-list-docx.py
#
# Generated from build/prices.js, the same file the PDF is built
# from. That is the whole point of it: converting the PDF to Word
# would give a document that looks roughly right and shares nothing
# with the source, so the next price change would have to be made
# twice and the two would drift the first time somebody forgot.
# Change a price in prices.js, re-run both, and they agree.
#
# What this is FOR: editing at the counter. Someone needs to add a
# line, move a decimal or send a customer a quotation with two
# products struck out. The PDF is the printed piece; this is the
# working copy.
#
# What it is NOT: the designed sheet. Word has no equivalent of the
# gradient ground and would make a poor job of pretending. The
# structure carries over — the categories in order, the tin colour on
# each product with its two lighter shades for pack and price — as
# real table shading, which Word can edit. It reads as the same
# document without lying about being the same object.
#
# Needs python-docx:  pip install python-docx
# ============================================================

import json
import os
import subprocess
import sys

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit('python-docx is not installed.  pip install python-docx')

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
OUT = os.path.join(PKG, 'pdf', '5-price-list', 'cloud-paints-price-list.docx')

BLUE = RGBColor(0x0F, 0x1F, 0x5C)
RED = RGBColor(0x8B, 0x1E, 0x2C)
INK = RGBColor(0x12, 0x14, 0x2B)
INK2 = RGBColor(0x4A, 0x4D, 0x68)


# ---- the range, resolved by node so the colours match the PDF exactly ----
NODE = r"""
import { PRICE_LIST, TIN_LABEL, CURRENCY, EFFECTIVE_FROM, TRADE_NOTE } from './build/prices.js';
import { loadProducts, readable, tint, inkOn } from './build/lib.mjs';
const bySlug = Object.fromEntries(loadProducts().map(p => [p.slug, p]));
const colourOf = row =>
  row.colour || TIN_LABEL[row.art]
  || (bySlug[row.art] ? readable(bySlug[row.art].primary) : '#4a5568');
console.log(JSON.stringify({
  currency: CURRENCY, effective: EFFECTIVE_FROM, note: TRADE_NOTE,
  groups: PRICE_LIST.map(g => ({
    title: g.title,
    accent: colourOf(g.rows[0]),
    rows: g.rows.map(r => {
      const c = colourOf(r);
      return { name: r.name, art: r.art, desc: r.desc,
               label: c, labelInk: inkOn(c),
               qty: tint(c, 0.74), qtyInk: inkOn(tint(c, 0.74)),
               pri: tint(c, 0.90), priInk: inkOn(tint(c, 0.90)),
               packs: Object.entries(r.prices) };
    }),
  })),
}));
"""


def load():
    out = subprocess.run(['node', '--input-type=module', '-e', NODE],
                         cwd=PKG, capture_output=True, text=True, encoding='utf-8')
    if out.returncode:
        sys.exit('could not read prices.js:\n' + out.stderr)
    return json.loads(out.stdout)


# ---- Word helpers -------------------------------------------------------
def shade(cell, hex_colour):
    """Solid cell fill. python-docx has no API for it, so the shading
    element goes on by hand."""
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:color'), 'auto')
    el.set(qn('w:fill'), hex_colour.lstrip('#'))
    cell._tc.get_or_add_tcPr().append(el)


def no_space(par):
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.space_after = Pt(0)


def put(cell, text, *, size=9, bold=False, colour=None, caps=False):
    cell.text = ''
    par = cell.paragraphs[0]
    no_space(par)
    run = par.add_run(text.upper() if caps else text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = 'Calibri'
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour.lstrip('#').upper())
    return par


def borders(table):
    """A visible grid. The PDF rules the cards heavily; the Word version
    should not arrive looking like a borderless list."""
    tbl = table._tbl.tblPr
    marks = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '8')
        e.set(qn('w:color'), '595959')
        marks.append(e)
    tbl.append(marks)


# ---- the document -------------------------------------------------------
def build():
    d = load()
    doc = Document()

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    for m in ('top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(1.4))
    for m in ('left_margin', 'right_margin'):
        setattr(sec, m, Cm(1.5))

    logo = os.path.join(PKG, 'assets', 'img', 'brand', 'logo.png')
    if os.path.exists(logo):
        p = doc.add_paragraph()
        no_space(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(logo, width=Cm(4.2))

    t = doc.add_paragraph()
    no_space(t)
    r = t.add_run('PRICE LIST')
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = BLUE
    r.font.name = 'Calibri'

    t2 = doc.add_paragraph()
    no_space(t2)
    r = t2.add_run('RECOMMENDED RETAIL PRICES')
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RED
    r.font.name = 'Calibri'

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run('All prices in Kenya Shillings'
                    + (' · Effective ' + d['effective'] if d['effective'] else ''))
    r.font.size = Pt(9)
    r.font.color.rgb = INK2
    r.font.name = 'Calibri'

    for g in d['groups']:
        # section plate — one full-width shaded row, as the PDF has
        head = doc.add_table(rows=1, cols=1)
        head.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = head.rows[0].cells[0]
        shade(c, g['accent'])
        ink = '#FFFFFF' if g['accent'] else '#FFFFFF'
        put(c, g['title'], size=12, bold=True, colour=ink, caps=True)
        head.rows[0].cells[0].paragraphs[0].paragraph_format.space_before = Pt(2)

        # one row per pack, product name merged down its packs
        total = sum(len(r['packs']) for r in g['rows'])
        tbl = doc.add_table(rows=total + 1, cols=3)
        borders(tbl)
        hdr = tbl.rows[0]
        for i, (txt, w) in enumerate((('PRODUCT', Cm(9.4)), ('PACK', Cm(3.4)),
                                      ('PRICE (KSHS)', Cm(4.6)))):
            put(hdr.cells[i], txt, size=8, bold=True, colour='#FFFFFF', caps=True)
            shade(hdr.cells[i], '#3F3F46')
            hdr.cells[i].width = w

        i = 1
        for row in g['rows']:
            first = i
            for pack, price in row['packs']:
                cells = tbl.rows[i].cells
                shade(cells[1], row['qty'])
                put(cells[1], pack, size=9, bold=True, colour=row['qtyInk'])
                shade(cells[2], row['pri'])
                put(cells[2],
                    '' if not price else '{:,}'.format(price),
                    size=11, bold=True, colour=row['priInk'])
                i += 1
            merged = tbl.cell(first, 0)
            if i - 1 > first:
                merged = merged.merge(tbl.cell(i - 1, 0))
            shade(merged, row['label'])
            put(merged, row['name'], size=10, bold=True,
                colour=row['labelInk'], caps=True)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ---- terms ----------------------------------------------------------
    h = doc.add_paragraph()
    r = h.add_run('TERMS AND CONDITIONS')
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = BLUE
    r.font.name = 'Calibri'

    finishes = ('Venetian Marble, Stone Texture, Luxury Stucco, Concrete Finish, '
                'Velvet Texture, Metallic Illusion, Rustic Texture, Sand Finish, '
                'Desert Stone and RockShield Exterior Texture')
    for text in (
        d['note'],
        'Colour Tinting is available at the Industrial Area Factory on all Emulsions '
        'and Enamels. Tinted Shades may carry a Surcharge depending on the Colourant used.',
        'Hand applied Decorative Finishes — ' + finishes + ' — are quoted separately. '
        'They are sold by Weight and are dependent on Wall Measurements. The Decorative '
        'Palette is available at our Factory and is priced per Quantity ordered.',
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = INK2
        r.font.name = 'Calibri'

    p = doc.add_paragraph()
    r = p.add_run('Manufactured by Cloudsent Decor Ltd · 10 Rangwe Road, off Lunga Lunga '
                  'Road, Industrial Area, Nairobi · P.O. Box 44192–00100\n'
                  'Call +254 741 405 481 · +254 788 866 620 · +254 727 779 085   '
                  'WhatsApp 0741 405 481 or 0727 779 085\n'
                  'info@cloudpaints.co.ke · www.cloudpaints.co.ke   '
                  'Instagram / Facebook / X: cloudpaintskenya')
    r.font.size = Pt(8.5)
    r.font.color.rgb = INK
    r.font.name = 'Calibri'

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)

    lines = sum(len(g['rows']) for g in d['groups'])
    packs = sum(len(r['packs']) for g in d['groups'] for r in g['rows'])
    print('  %s' % OUT)
    print('  %d categories · %d lines · %d priced packs · %.0f KB'
          % (len(d['groups']), lines, packs, os.path.getsize(OUT) / 1024))
    print('')
    print('  Generated from build/prices.js — the same file the PDF is built from.')
    print('  Change a price there and re-run both; edit this file directly and it')
    print('  becomes a one-off that the next build will not know about.')


if __name__ == '__main__':
    build()
